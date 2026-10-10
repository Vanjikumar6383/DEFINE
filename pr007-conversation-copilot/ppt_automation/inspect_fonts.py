from pptx import Presentation

pptx_path = r"C:\Users\psamy\define\pr007-conversation-copilot\General - DEFINE Template.pptx"
prs = Presentation(pptx_path)

for s_idx in [0, 1]:
    slide = prs.slides[s_idx]
    print(f"\n--- Slide {s_idx+1} Details ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    color = None
                    if r.font.color and r.font.color.type:
                        try:
                            color = r.font.color.rgb
                        except:
                            color = "theme"
                    print(f"Shape: {shape.name} | Text: '{r.text}' | Font: {r.font.name} | Size: {r.font.size} | Color: {color} | Bold: {r.font.bold}")
