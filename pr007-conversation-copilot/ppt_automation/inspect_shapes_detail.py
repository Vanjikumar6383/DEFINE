import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r"C:\Users\psamy\define\pr007-conversation-copilot\General - DEFINE Template.pptx"
prs = Presentation(pptx_path)

def emu_to_in(emu):
    return round(emu / 914400, 3)

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        t_text = ""
        font_info = ""
        if shape.has_text_frame:
            paras = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            t_text = " // ".join(paras)
            if shape.text_frame.paragraphs:
                p0 = shape.text_frame.paragraphs[0]
                if p0.runs:
                    r0 = p0.runs[0]
                    font_info = f"font: {r0.font.name}, size: {r0.font.size}, bold: {r0.font.bold}"
        print(f"Shape {s_idx} [{shape.name}] type: {shape.shape_type} | pos: ({emu_to_in(shape.left)}\", {emu_to_in(shape.top)}\") size: ({emu_to_in(shape.width)}\" x {emu_to_in(shape.height)}\") | text: '{t_text[:80]}' | {font_info}")
