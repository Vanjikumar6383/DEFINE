import os
import sys
from pptx import Presentation

pptx_path = r"C:\Users\psamy\define\pr007-conversation-copilot\General - DEFINE Template.pptx"
prs = Presentation(pptx_path)

print(f"Total slides: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width}, height: {prs.slide_height}")

for idx, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {idx + 1} ---")
    print(f"Layout name: {slide.slide_layout.name}")
    shapes_info = []
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        text = ""
        if shape.has_text_frame:
            text = " | ".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        elif shape.has_table:
            text = f"TABLE {len(shape.table.rows)}x{len(shape.table.columns)}"
        shapes_info.append(f"  Shape {s_idx}: {shape.name} (type: {shape_type}) -> text: {text[:150]}")
    for info in shapes_info:
        print(info)
