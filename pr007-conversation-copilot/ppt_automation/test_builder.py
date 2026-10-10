from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

pptx_path = r"C:\Users\psamy\define\pr007-conversation-copilot\General - DEFINE Template.pptx"
prs = Presentation(pptx_path)

print(f"Loaded template with {len(prs.slides)} slides.")
print("Verifying python-pptx capabilities...")
